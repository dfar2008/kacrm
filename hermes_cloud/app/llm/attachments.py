"""Attachment ingestion for the AI assistants — shared by the schema
builder assistant and the dashboard assistant.

Two concerns:
  • extract_text(filename, data) — pull plain text out of an uploaded
    document (PDF / Excel / Word / CSV / plain text).
  • prepare_attachments(files) — classify a batch of uploads into a
    normalized shape both assistants' build_messages() consume:
      image    → {"kind": "image",    "filename": str, "data_url": str}
      document → {"kind": "document", "filename": str, "text": str}

Images are base64 data URLs (fed to vision-capable LLMs as image_url
content blocks). Documents are extracted to text server-side, so they
work with any text LLM.
"""
from __future__ import annotations

import base64
import csv
import io
import mimetypes
from typing import Any, Dict, List, Optional, Tuple

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}

# Cap extracted text per document so one huge file can't blow the LLM
# context window. 20k chars ≈ a few thousand tokens.
_MAX_DOC_CHARS = 20000


def _ext(filename: str) -> str:
    i = (filename or "").rfind(".")
    return filename[i:].lower() if i >= 0 else ""


def _decode(data: bytes) -> str:
    """Best-effort decode of a text payload — tries UTF-8, then GB18030
    (Chinese), then latin-1, finally a lossy UTF-8."""
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: List[str] = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append("\t".join(cells))
    return "\n".join(out)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: List[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(_decode(data))))
    return "\n".join("\t".join(r) for r in rows)


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a document upload, dispatched by extension.

    Unknown extensions fall back to a best-effort text decode.
    """
    ext = _ext(filename)
    if ext == ".pdf":
        return _from_pdf(data)
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(data)
    if ext == ".docx":
        return _from_docx(data)
    if ext in (".csv", ".tsv"):
        return _from_csv(data)
    return _decode(data)


# Below this many characters of extracted text, a PDF is treated as
# scanned / image-only (no usable text layer) and rendered to page images so
# the vision model can still read it — e.g. a photographed invoice saved as PDF.
_PDF_SCANNED_TEXT_MIN = 40
_PDF_MAX_RENDER_PAGES = 8


def _pdf_to_image_dataurls(
    data: bytes, max_pages: int = _PDF_MAX_RENDER_PAGES,
) -> List[str]:
    """Render PDF pages to PNG data URLs for vision fallback on scanned PDFs.

    Returns [] if rendering is unavailable or fails, so callers degrade to the
    (empty) text path instead of erroring.
    """
    try:
        import pypdfium2 as pdfium  # BSD-licensed, bundles PDFium
    except Exception:
        return []
    urls: List[str] = []
    try:
        pdf = pdfium.PdfDocument(data)
        try:
            for i in range(min(len(pdf), max_pages)):
                bitmap = pdf[i].render(scale=2.0)  # ~144 dpi — legible for OCR/vision
                buf = io.BytesIO()
                bitmap.to_pil().save(buf, format="PNG")
                b64 = base64.b64encode(buf.getvalue()).decode("ascii")
                urls.append(f"data:image/png;base64,{b64}")
        finally:
            pdf.close()
    except Exception:
        return []
    return urls


def build_user_content(
    user_text: str,
    attachments: Optional[List[Dict[str, Any]]] = None,
):
    """Compose an LLM user-message ``content`` from text + attachments.

    Documents (extracted text) are appended to the prompt text. Images
    become ``image_url`` content blocks — the message content becomes a
    multimodal list. With no images the content stays a plain string so
    non-vision providers are unaffected.
    """
    attachments = attachments or []
    docs = [a for a in attachments if a.get("kind") == "document"]
    images = [a for a in attachments if a.get("kind") == "image"]
    text = user_text
    for d in docs:
        text += f"\n\n[附件:{d.get('filename', 'file')}]\n{d.get('text', '')}"
    if not images:
        return text
    return [
        {"type": "text", "text": text},
        *[
            {"type": "image_url", "image_url": {"url": a["data_url"]}}
            for a in images
        ],
    ]


def prepare_attachments(
    files: List[Tuple[str, bytes, str]],
) -> List[Dict[str, Any]]:
    """Classify a batch of uploads for an assistant prompt.

    ``files`` items are ``(filename, data, content_type)``. Returns a
    normalized list — images carry a base64 data URL, documents carry
    extracted (and length-capped) text.
    """
    out: List[Dict[str, Any]] = []
    for filename, data, content_type in files:
        ext = _ext(filename)
        is_image = ext in IMAGE_EXTS or (content_type or "").startswith("image/")
        if is_image:
            mime = (
                content_type
                or mimetypes.guess_type(filename)[0]
                or "image/png"
            )
            b64 = base64.b64encode(data).decode("ascii")
            out.append({
                "kind": "image",
                "filename": filename,
                "data_url": f"data:{mime};base64,{b64}",
            })
        else:
            text = extract_text(filename, data)
            # Scanned / image-only PDF (no text layer): render its pages to
            # images so the vision model can read it. Without this, a
            # photographed invoice saved as PDF yields ~empty text and the
            # agent sees nothing (while the same invoice sent as a JPG works).
            if ext == ".pdf" and len(text.strip()) < _PDF_SCANNED_TEXT_MIN:
                page_urls = _pdf_to_image_dataurls(data)
                if page_urls:
                    for j, url in enumerate(page_urls):
                        out.append({
                            "kind": "image",
                            "filename": f"{filename} (p{j + 1})",
                            "data_url": url,
                        })
                    continue
            if len(text) > _MAX_DOC_CHARS:
                text = text[:_MAX_DOC_CHARS] + "\n…(truncated)"
            out.append({
                "kind": "document",
                "filename": filename,
                "text": text,
            })
    return out
